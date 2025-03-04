# converter.py
import os
import shutil
import sqlite3
import fitz  # PyMuPDF
from pdf2docx import Converter as PDF2DocxConverter
from docx import Document
import pytesseract
from PIL import Image
import tempfile
import datetime

# Optional: For Word to PDF conversion using docx2pdf.
try:
    from docx2pdf import convert as docx2pdf_convert
except ImportError:
    docx2pdf_convert = None

DB_PATH = "history.db"

def add_history(file_name, conversion_type):
    """Insert a new conversion record into the history database."""
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO history (file_name, conversion_type) VALUES (?, ?)",
        (file_name, conversion_type)
    )
    conn.commit()
    conn.close()

def extract_text_with_ocr(pdf_path):
    """
    Process a scanned PDF by converting each page to an image and running OCR.
    Returns the concatenated text from all pages.
    """
    doc = fitz.open(pdf_path)
    full_text = ""
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        pix = page.get_pixmap()  # Render page to image.
        # Save the image temporarily.
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(pix.tobytes("png"))
            tmp_path = tmp.name
        # Run OCR on the saved image.
        text = pytesseract.image_to_string(Image.open(tmp_path))
        full_text += text + "\n"
        os.remove(tmp_path)
    return full_text

def convert_pdf_to_word(input_pdf, output_docx):
    """
    Convert a PDF file to a Word document.
    If the PDF appears to be scanned (i.e., no text), use OCR.
    """
    # Attempt to extract text using pdf2docx.
    try:
        cv = PDF2DocxConverter(input_pdf)
        cv.convert(output_docx, start=0, end=None)
        cv.close()
        # Check if the resulting document has content.
        doc = Document(output_docx)
        if not any(para.text.strip() for para in doc.paragraphs):
            raise ValueError("Empty document - possibly a scanned PDF.")
    except Exception as e:
        # Fallback to OCR conversion.
        print("Standard conversion failed or PDF appears scanned. Using OCR.")
        text = extract_text_with_ocr(input_pdf)
        # Create a new Word document with the OCR text.
        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)
        doc.save(output_docx)
    
    add_history(output_docx, "PDF to Word")
    return output_docx

def convert_word_to_pdf(input_docx, output_pdf):
    """
    Convert a Word document to PDF.
    Uses docx2pdf if available; otherwise, a basic conversion is performed.
    """
    if docx2pdf_convert:
        try:
            # docx2pdf expects file paths.
            docx2pdf_convert(input_docx, output_pdf)
        except Exception as e:
            raise RuntimeError("Error during Word to PDF conversion: " + str(e))
    else:
        # Fallback: Generate a PDF from text (limited formatting).
        doc = Document(input_docx)
        pdf_text = "\n".join([para.text for para in doc.paragraphs])
        # Create a basic PDF using PyMuPDF.
        pdf_doc = fitz.open()
        page = pdf_doc.new_page()
        page.insert_text((72, 72), pdf_text)
        pdf_doc.save(output_pdf)
        pdf_doc.close()
    
    add_history(output_pdf, "Word to PDF")
    return output_pdf
